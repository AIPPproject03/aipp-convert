import os
import sys
import time
import shutil
import uuid
import threading
import subprocess
from pathlib import Path
from datetime import datetime, timedelta
from flask import Flask, request, jsonify, send_file, render_template
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

# Memuat konfigurasi dari .env jika tersedia
load_dotenv()

# Inisialisasi Flask
app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 200 * 1024 * 1024  # Batas upload 200 MB (dapat disesuaikan)

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_FOLDER = BASE_DIR / 'storage' / 'uploads'
CONVERTED_FOLDER = BASE_DIR / 'storage' / 'converted'

UPLOAD_FOLDER.mkdir(parents=True, exist_ok=True)
CONVERTED_FOLDER.mkdir(parents=True, exist_ok=True)

ALLOWED_EXTENSIONS_PDF = {'pdf'}
ALLOWED_EXTENSIONS_WORD = {'docx', 'doc'}

# ==============================================================================
# KONFIGURASI TESSERACT OCR & POPPLER (OTOMATIS & LINTAS PLATFORM)
# ==============================================================================
import pytesseract
from pdf2image import convert_from_path
import docx
from pypdf import PdfReader
from pdf2docx import Converter

def configure_ocr_and_poppler():
    """Mendeteksi instalasi Tesseract dan Poppler secara otomatis di Windows / Linux / macOS."""
    poppler_path = os.getenv('POPPLER_PATH')
    tesseract_cmd = os.getenv('TESSERACT_CMD')

    # Konfigurasi Tesseract CMD jika di Windows
    if sys.platform.startswith('win'):
        if tesseract_cmd and os.path.exists(tesseract_cmd):
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
        else:
            common_tesseract_paths = [
                r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe")
            ]
            for p in common_tesseract_paths:
                if os.path.exists(p):
                    pytesseract.pytesseract.tesseract_cmd = p
                    break

        # Cek Poppler default jika belum didefinisikan
        if not poppler_path:
            common_poppler_paths = [
                r"C:\poppler\Library\bin",
                r"C:\Program Files\poppler\Library\bin",
                r"C:\Program Files\poppler\bin",
                os.path.expandvars(r"%LOCALAPPDATA%\poppler\Library\bin")
            ]
            for p in common_poppler_paths:
                if os.path.exists(p):
                    poppler_path = p
                    break

    return poppler_path

POPPLER_PATH = configure_ocr_and_poppler()

# ==============================================================================
# SISTEM DETEKSI KESEHATAN DEPENDENSI
# ==============================================================================
def check_system_capabilities():
    """Memeriksa status dependensi pendukung untuk ditampilkan di antarmuka."""
    status = {
        'pdf2docx': True,
        'tesseract_installed': False,
        'tesseract_path': None,
        'poppler_installed': False,
        'poppler_path': POPPLER_PATH,
        'word_converter': 'none'
    }

    # Cek Tesseract
    try:
        tess_version = pytesseract.get_tesseract_version()
        status['tesseract_installed'] = True
        status['tesseract_version'] = str(tess_version)
        status['tesseract_path'] = pytesseract.pytesseract.tesseract_cmd
    except Exception:
        status['tesseract_installed'] = False

    # Cek Poppler
    try:
        if POPPLER_PATH and os.path.exists(POPPLER_PATH):
            status['poppler_installed'] = True
        else:
            # Uji apakah pdftoppm ada di PATH
            which_res = shutil.which('pdftoppm')
            if which_res:
                status['poppler_installed'] = True
                status['poppler_path'] = os.path.dirname(which_res)
    except Exception:
        status['poppler_installed'] = False

    # Cek Mesin Word -> PDF (docx2pdf via Word COM / LibreOffice)
    if sys.platform.startswith('win'):
        status['word_converter'] = 'docx2pdf (MS Word / COM)'
    elif shutil.which('soffice') or shutil.which('libreoffice'):
        status['word_converter'] = 'LibreOffice (Headless)'
    else:
        status['word_converter'] = 'docx2pdf / LibreOffice'

    return status

# ==============================================================================
# UTILITAS VALIDASI & MANAJEMEN BERKAS
# ==============================================================================
def allowed_file(filename, allowed_set):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in allowed_set

def generate_unique_filename(original_filename, target_extension):
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    unique_id = uuid.uuid4().hex[:6]
    base_name = Path(original_filename).stem
    safe_base = "".join(c for c in base_name if c.isalnum() or c in (' ', '_', '-')).rstrip()
    if not safe_base:
        safe_base = "dokumen"
    return f"{safe_base}_{timestamp}_{unique_id}.{target_extension}"

# ==============================================================================
# LOGIKA PEMBERSIHAN BERKAS OTOMATIS (BACKGROUND THREAD)
# ==============================================================================
def background_cleanup_job(max_age_seconds=3600, interval_seconds=900):
    """
    Menghapus berkas sementara yang berumur lebih dari 1 jam (3600 detik).
    Dijalankan secara berkala di latar belakang setiap 15 menit.
    """
    while True:
        try:
            now = time.time()
            for folder in [UPLOAD_FOLDER, CONVERTED_FOLDER]:
                for item in folder.glob('*'):
                    if item.is_file():
                        file_age = now - item.stat().st_mtime
                        if file_age > max_age_seconds:
                            try:
                                item.unlink()
                                print(f"[Cleanup] Menghapus berkas kedaluwarsa: {item.name}")
                            except Exception as e:
                                print(f"[Cleanup Error] Gagal menghapus {item.name}: {e}")
        except Exception as e:
            print(f"[Cleanup Background Exception]: {e}")
        time.sleep(interval_seconds)

# Jalankan daemon thread untuk pembersihan berkas
cleanup_thread = threading.Thread(target=background_cleanup_job, daemon=True)
cleanup_thread.start()

# ==============================================================================
# LOGIKA DETEKSI LAPISAN TEKS vs SCAN (OCR)
# ==============================================================================
def detect_pdf_type(pdf_path):
    """
    Menganalisis apakah PDF memiliki lapisan teks yang dapat diekstrak
    atau merupakan hasil scan (berbasis gambar murni).
    Mengembalikan: (is_text_based: bool, total_chars: int, page_count: int)
    """
    try:
        reader = PdfReader(pdf_path)
        total_pages = len(reader.pages)
        if total_pages == 0:
            return False, 0, 0

        # Cek hingga 5 halaman pertama untuk sampel deteksi
        sample_pages = min(total_pages, 5)
        extracted_chars = 0

        for i in range(sample_pages):
            page_text = reader.pages[i].extract_text() or ""
            # Bersihkan spasi kosong berlebih
            clean_text = "".join(page_text.split())
            extracted_chars += len(clean_text)

        # Rata-rata karakter per halaman sampel
        avg_chars = extracted_chars / sample_pages

        # Jika rata-rata karakter per halaman lebih dari 40 karakter,
        # dokumen dikategorikan sebagai PDF Berbasis Teks asli.
        is_text_based = avg_chars >= 40
        return is_text_based, extracted_chars, total_pages
    except Exception as e:
        print(f"[Deteksi PDF Error] {e}")
        # Asumsi aman jika pembacaan teks gagal: perlakukan sebagai scan
        return False, 0, 0

# ==============================================================================
# METODE KONVERSI PDF -> WORD
# ==============================================================================
def convert_pdf_to_word_layout(pdf_path, output_docx_path):
    """Metode 1: Ekstraksi langsung dengan pdf2docx untuk mempertahankan layout & tabel."""
    cv = Converter(str(pdf_path))
    cv.convert(str(output_docx_path), start=0, end=None)
    cv.close()

def convert_pdf_to_word_ocr(pdf_path, output_docx_path, ocr_lang='ind+eng'):
    """Metode 2: OCR menggunakan pdf2image + pytesseract + python-docx untuk PDF scan/foto."""
    kwargs = {}
    if POPPLER_PATH:
        kwargs['poppler_path'] = POPPLER_PATH

    # Konversi setiap halaman PDF menjadi gambar PIL
    images = convert_from_path(str(pdf_path), dpi=200, **kwargs)
    
    doc = docx.Document()

    for idx, img in enumerate(images, start=1):
        # Ekstraksi teks dari gambar halaman
        text = pytesseract.image_to_string(img, lang=ocr_lang)

        # Tambahkan judul halaman jika lebih dari 1 halaman
        if len(images) > 1:
            p_hdr = doc.add_paragraph()
            run_hdr = p_hdr.add_run(f"--- [Halaman {idx}] ---")
            run_hdr.italic = True
            run_hdr.font.size = docx.shared.Pt(9)
            run_hdr.font.color.rgb = docx.shared.RGBColor(128, 128, 128)

        # Masukkan teks per paragraf
        paragraphs = text.split('\n\n')
        for p in paragraphs:
            clean_p = p.strip()
            if clean_p:
                doc.add_paragraph(clean_p)

        # Pemisah halaman jika bukan halaman terakhir
        if idx < len(images):
            doc.add_page_break()

    doc.save(str(output_docx_path))

# ==============================================================================
# METODE KONVERSI WORD -> PDF
# ==============================================================================
def convert_word_to_pdf_engine(docx_path, output_pdf_path):
    """
    Mengonversi dokumen Word ke PDF dengan pendekatan multi-engine:
    1. Mencoba `docx2pdf` (sangat optimal jika di Windows dengan MS Word)
    2. Fallback ke LibreOffice Headless (berfungsi di Windows/Linux/macOS)
    """
    error_log = []

    # Opsi 1: Coba docx2pdf
    try:
        from docx2pdf import convert
        convert(str(docx_path), str(output_pdf_path))
        if os.path.exists(output_pdf_path) and os.path.getsize(output_pdf_path) > 0:
            return True, "docx2pdf"
    except Exception as e:
        error_log.append(f"docx2pdf failed: {str(e)}")

    # Opsi 2: Fallback ke LibreOffice / soffice
    libreoffice_bins = ['soffice', 'libreoffice', r"C:\Program Files\LibreOffice\program\soffice.exe"]
    for bin_name in libreoffice_bins:
        if shutil.which(bin_name) or (os.path.exists(bin_name) if os.path.isabs(bin_name) else False):
            try:
                cmd = [
                    bin_name,
                    '--headless',
                    '--convert-to', 'pdf',
                    '--outdir', str(CONVERTED_FOLDER),
                    str(docx_path)
                ]
                subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=120)
                # LibreOffice menghasilkan output dengan nama file sama di outdir
                expected_lo_output = CONVERTED_FOLDER / f"{docx_path.stem}.pdf"
                if expected_lo_output.exists():
                    if expected_lo_output != output_pdf_path:
                        shutil.move(str(expected_lo_output), str(output_pdf_path))
                    return True, "LibreOffice"
            except Exception as e:
                error_log.append(f"LibreOffice failed: {str(e)}")

    raise RuntimeError(
        "Gagal mengonversi Word ke PDF. Pastikan Microsoft Word (di Windows) "
        "atau LibreOffice terinstal di sistem untuk konversi DOCX ke PDF. "
        + " | ".join(error_log)
    )

# ==============================================================================
# ROUTE FLASK
# ==============================================================================
@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/system-status', methods=['GET'])
def system_status():
    """Mengembalikan status kesiapan pustaka & tools konversi."""
    capabilities = check_system_capabilities()
    return jsonify({
        'success': True,
        'capabilities': capabilities
    })

@app.route('/api/convert/pdf-to-word', methods=['POST'])
def api_pdf_to_word():
    """Endpoint konversi PDF ke Word dengan deteksi otomatis atau mode OCR manual."""
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'Tidak ada berkas yang diunggah.'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'message': 'Nama berkas kosong.'}), 400

    if not allowed_file(file.filename, ALLOWED_EXTENSIONS_PDF):
        return jsonify({'success': False, 'message': 'Format berkas tidak didukung. Mohon unggah berkas .pdf.'}), 400

    force_ocr = request.form.get('force_ocr', 'false').lower() == 'true'
    ocr_lang = request.form.get('ocr_lang', 'ind+eng')

    # Buat nama berkas unik
    original_name = secure_filename(file.filename) or "dokumen.pdf"
    input_filename = f"input_{uuid.uuid4().hex[:8]}_{original_name}"
    input_path = UPLOAD_FOLDER / input_filename
    file.save(str(input_path))

    output_filename = generate_unique_filename(original_name, 'docx')
    output_path = CONVERTED_FOLDER / output_filename

    try:
        is_text_based, char_count, total_pages = detect_pdf_type(input_path)
        method_used = ""
        detection_label = "Teks Terstruktur" if is_text_based else "Gambar / Scan"

        if force_ocr or not is_text_based:
            # Mode OCR
            method_used = "OCR (Tesseract)"
            try:
                convert_pdf_to_word_ocr(input_path, output_path, ocr_lang=ocr_lang)
            except Exception as ocr_err:
                # Jika OCR gagal karena Tesseract / Poppler belum disiapkan, berikan pesan panduan
                err_msg = str(ocr_err)
                if "tesseract is not installed" in err_msg.lower() or "tesseract" in err_msg.lower():
                    raise RuntimeError(
                        "Tesseract OCR belum terpasang atau belum terbaca di PATH sistem. "
                        "Silakan pasang Tesseract-OCR di server/komputer Anda."
                    )
                elif "poppler" in err_msg.lower() or "unable to get page count" in err_msg.lower():
                    raise RuntimeError(
                        "Poppler belum terpasang atau jalurnya belum terkonfigurasi. "
                        "Poppler diperlukan untuk merender halaman PDF ke gambar sebelum OCR."
                    )
                else:
                    raise ocr_err
        else:
            # Mode Tata Letak Asli (pdf2docx)
            method_used = "Ekstraksi Layout Asli (pdf2docx)"
            try:
                convert_pdf_to_word_layout(input_path, output_path)
            except Exception as layout_err:
                # Jika layout gagal, coba fallback ke OCR jika dimungkinkan
                print(f"[Layout Fail -> Fallback OCR]: {layout_err}")
                try:
                    convert_pdf_to_word_ocr(input_path, output_path, ocr_lang=ocr_lang)
                    method_used = "OCR Fallback (Tesseract)"
                except Exception:
                    raise layout_err

        file_size_bytes = output_path.stat().st_size
        file_size_formatted = f"{file_size_bytes / 1024:.1f} KB" if file_size_bytes < 1024 * 1024 else f"{file_size_bytes / (1024*1024):.2f} MB"

        return jsonify({
            'success': True,
            'message': 'Konversi PDF ke Word berhasil!',
            'download_url': f'/api/download/converted/{output_filename}',
            'filename': output_filename,
            'original_filename': original_name,
            'method_used': method_used,
            'detected_type': detection_label,
            'total_pages': total_pages,
            'file_size': file_size_formatted
        })

    except Exception as e:
        print(f"[Error PDF to Word]: {e}")
        return jsonify({
            'success': False,
            'message': f"Terjadi kesalahan saat konversi: {str(e)}"
        }), 500

    finally:
        # Hapus berkas upload input sementara untuk efisiensi penyimpanan
        if input_path.exists():
            try:
                input_path.unlink()
            except Exception:
                pass

@app.route('/api/convert/word-to-pdf', methods=['POST'])
def api_word_to_pdf():
    """Endpoint konversi Word (.docx / .doc) ke PDF."""
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'Tidak ada berkas yang diunggah.'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'message': 'Nama berkas kosong.'}), 400

    if not allowed_file(file.filename, ALLOWED_EXTENSIONS_WORD):
        return jsonify({'success': False, 'message': 'Format berkas tidak didukung. Unggah berkas .docx atau .doc.'}), 400

    original_name = secure_filename(file.filename) or "dokumen.docx"
    input_filename = f"input_{uuid.uuid4().hex[:8]}_{original_name}"
    input_path = UPLOAD_FOLDER / input_filename
    file.save(str(input_path))

    output_filename = generate_unique_filename(original_name, 'pdf')
    output_path = CONVERTED_FOLDER / output_filename

    try:
        success, engine_name = convert_word_to_pdf_engine(input_path, output_path)

        file_size_bytes = output_path.stat().st_size
        file_size_formatted = f"{file_size_bytes / 1024:.1f} KB" if file_size_bytes < 1024 * 1024 else f"{file_size_bytes / (1024*1024):.2f} MB"

        return jsonify({
            'success': True,
            'message': 'Konversi Word ke PDF berhasil!',
            'download_url': f'/api/download/converted/{output_filename}',
            'filename': output_filename,
            'original_filename': original_name,
            'method_used': f"Engine {engine_name}",
            'file_size': file_size_formatted
        })

    except Exception as e:
        print(f"[Error Word to PDF]: {e}")
        return jsonify({
            'success': False,
            'message': f"Gagal mengonversi Word ke PDF: {str(e)}"
        }), 500

    finally:
        # Hapus berkas upload input sementara
        if input_path.exists():
            try:
                input_path.unlink()
            except Exception:
                pass

@app.route('/api/download/<folder_type>/<filename>', methods=['GET'])
def download_file(folder_type, filename):
    """Endpoint pengunduhan berkas hasil konversi."""
    safe_filename = secure_filename(filename)
    if folder_type == 'converted':
        target_path = CONVERTED_FOLDER / safe_filename
    else:
        target_path = UPLOAD_FOLDER / safe_filename

    if not target_path.exists():
        return jsonify({'error': 'Berkas tidak ditemukan atau telah dibersihkan otomatis.'}), 404

    return send_file(
        str(target_path),
        as_attachment=True,
        download_name=safe_filename
    )

# ==============================================================================
# JALANKAN SERVER
# ==============================================================================
if __name__ == '__main__':
    port = int(os.getenv('PORT', 5000))
    print("=" * 60)
    print("  🚀 aipp-convert: Aplikasi Konversi Dokumen Mandiri & Privat")
    print(f"  🌐 URL Lokal: http://127.0.0.1:{port}")
    print("=" * 60)
    app.run(host='0.0.0.0', port=port, debug=True)
